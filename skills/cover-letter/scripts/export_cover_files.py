"""export_cover_files.py : render a cover-letter .md into upload-ready PDF + Word files.

The drafted cover letter is stored as markdown (output/cover_letters/..._<jobid>.md), which is
the tracked source registered on the job. But many application forms require the cover letter as
an uploaded FILE, not pasted text, so this renders that letter into a clean one-page document in
BOTH formats:
  - <Name>-CoverLetter-<Company>.pdf   (primary: works on most upload fields; text-based,
    fonts embedded so accented characters always render)
  - <Name>-CoverLetter-<Company>.docx  (for Word-only forms and recruitment agencies)

The rendered document adds a simple name + contact header (so the file stands alone), then the
letter body. The job id is stripped from the file name for a clean send-ready name; when a company
has more than one letter, pass --role-slug (same rule as the CV PDFs) so they do not collide.

Contact details come from config/application_rules.md: pass --market <code> (a market from
config/filters.md, for example uk) to pick the matching phone number (--phone and --email
override). The page is A4. Requires python-docx
and LibreOffice (soffice).

Usage:
  python3 skills/cover-letter/scripts/export_cover_files.py \
      --md output/cover_letters/acme_product-marketing-manager_12.md \
      --out-base Jordan-Bennett-CoverLetter-Acme --market uk
  # a company with more than one letter:
  ... --out-base Jordan-Bennett-CoverLetter-Acme --role-slug DataAnalyst --market uk
Prints JSON {md, docx, pdf, email, phone}.
"""
import json
import argparse
import subprocess
import tempfile
import shutil
from pathlib import Path

import sys

import docx
from docx.shared import Pt, Mm, RGBColor

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT))
from lib import config_loader  # noqa: E402

def contact_defaults():
    """Email, LinkedIn and per-market phones from config/application_rules.md."""
    return config_loader.contact()


def render_docx(md_path, out_docx, phone, email, linkedin, name=None):
    name = name or config_loader.candidate()["name"]
    text = Path(md_path).read_text(encoding="utf-8").strip()
    blocks = [b.strip() for b in text.split("\n\n") if b.strip()]
    d = docx.Document()
    s = d.sections[0]
    s.page_width = Mm(210); s.page_height = Mm(297)  # A4, like the CVs
    s.top_margin = Mm(20); s.bottom_margin = Mm(20); s.left_margin = Mm(22); s.right_margin = Mm(22)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    h = d.add_paragraph(); r = h.add_run(name); r.bold = True; r.font.size = Pt(16)
    h.paragraph_format.space_after = Pt(2)
    c = d.add_paragraph(); rc = c.add_run(f"{phone}  |  {email}  |  {linkedin}")
    rc.font.size = Pt(10); rc.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    c.paragraph_format.space_after = Pt(12)
    for bl in blocks:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.08
        for i, ln in enumerate(bl.split("\n")):
            if i:
                p.add_run().add_break()
            p.add_run(ln)
    d.save(out_docx)


def to_pdf(docx_path, outdir, soffice="soffice"):
    with tempfile.TemporaryDirectory() as td:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", td, str(docx_path)],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        produced = Path(td) / (Path(docx_path).stem + ".pdf")
        target = Path(outdir) / produced.name
        # The project folder is a mount, so the temp dir is on another filesystem and
        # Path.replace raises OSError 18 (invalid cross-device link). shutil.move copies
        # across devices, then removes the source.
        shutil.move(str(produced), str(target))
        return target


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True, help="path to the drafted cover-letter .md")
    ap.add_argument("--out-base", required=True,
                    help="clean file stem, e.g. Jordan-Bennett-CoverLetter-Acme (no id, no extension)")
    ap.add_argument("--role-slug", default=None, help="disambiguator when a company has >1 letter")
    ap.add_argument("--market", default=None,
                    help="market code (e.g. uk); picks the phone number from config/application_rules.md")
    ap.add_argument("--phone", default=None, help="override the phone number")
    ap.add_argument("--email", default=None, help="override the email address")
    ap.add_argument("--outdir", default=None, help="default: the .md's own folder")
    ap.add_argument("--soffice", default="soffice")
    a = ap.parse_args(argv)
    outdir = Path(a.outdir) if a.outdir else Path(a.md).resolve().parent
    outdir.mkdir(parents=True, exist_ok=True)
    base = a.out_base + (f"-{a.role_slug.strip('-')}" if a.role_slug else "")
    out_docx = outdir / (base + ".docx")
    c = contact_defaults()
    if not a.phone and not a.market:
        ap.error("pass --market <code> (or --phone) so the letter carries the right phone number")
    phone = a.phone or c["phones"].get(str(a.market).lower())
    if not phone:
        ap.error(f"no phone for market '{a.market}' in config/application_rules.md contact.phones; "
                 "add one or pass --phone")
    email = a.email or c["email"]
    render_docx(a.md, str(out_docx), phone, email, c["linkedin"] or c["website"])
    out_pdf = to_pdf(str(out_docx), str(outdir), a.soffice)
    print(json.dumps({"md": a.md, "docx": str(out_docx), "pdf": str(out_pdf),
                      "email": email, "phone": phone}, ensure_ascii=False))


if __name__ == "__main__":
    main()

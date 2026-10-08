#!/usr/bin/env python3
"""
check_grounding.py - Catch CV content that the MASTER PROFILE does not support.

The cardinal rule of this skill: never put anything on a CV that the candidate's
master profile doesn't genuinely back up. A model can still "extrapolate" - turn
"operations & strategy" into "procurement and supplier negotiation", or invent a
metric - and prose guidance alone won't always stop it. This script is the
mechanical safety net. It compares the finished CV against the master profile and
flags:

  * NUMBERS / METRICS in the CV that do not appear in the master profile
    (a fabricated figure is the most damaging error a CV can carry), and
  * CONTENT WORDS in the CV whose stem/prefix is not found in the master profile
    - i.e. claims, skills, tools, or activities that may have been inferred or
    invented - showing the exact line each appears on.

It deliberately OVER-flags: a legitimate rephrasing that mirrors a job ad
("client liaison" -> "stakeholder management") shows up here too. That is the
point - a human (or the model) then confirms each flagged line is a fair
restatement of something genuinely in the profile, and CUTS whatever isn't.
Zero flags is not the goal; a reviewed, honest CV is. Identity lines (name,
target headline), section headings, and contact details are ignored.

Usage:
    python check_grounding.py cv.docx master_profile.(md|txt|docx) [--no-color]
"""
import re
import sys
from docx import Document

_C = sys.stdout.isatty() and "--no-color" not in sys.argv
GREEN = "\033[92m" if _C else ""
RED = "\033[91m" if _C else ""
YELLOW = "\033[93m" if _C else ""
RESET = "\033[0m" if _C else ""

STOP = set("""
a an the and or but nor if then else of to in on at for with by from as into across through over
under between within per via up out off than that this these those it itself its their our your my
we they them he she you i is are was were be been being am do does did doing done has have had will
would can could shall should may might must not no yes also both each all any some more most much
many other another such only own same so too very just about after before while since until upon
including include included provides provide provided providing support supported supporting helps
help helped ensure ensuring ensured deliver delivers delivered delivering using used use new key
role roles team teams work works working worked experience year years month months day days daily
weekly monthly quarterly annual onto toward towards plus etc level levels well able first second
third strong proven track record throughout responsible duties various several based present current
ongoing including
""".split())

HEADINGS = {
    "professional summary", "summary", "profile", "personal profile", "personal statement",
    "key skills", "core skills", "skills", "technical skills",
    "work experience", "professional experience", "experience", "employment history", "career history",
    "education", "qualifications", "education and qualifications",
    "certifications", "certification", "professional development", "training",
    "languages", "interests", "hobbies", "references", "projects", "publications",
    "awards", "memberships", "professional memberships", "key achievements", "achievements",
}


def read_text(path):
    if path.lower().endswith(".docx"):
        return "\n".join(p.text for p in Document(path).paragraphs)
    return open(path, encoding="utf-8", errors="ignore").read()


def get_lines(path):
    if path.lower().endswith(".docx"):
        return [p.text for p in Document(path).paragraphs if p.text.strip()]
    return [l for l in read_text(path).split("\n") if l.strip()]


def is_contact(line):
    l = line.lower()
    return ("@" in l) or ("linkedin" in l) or ("http" in l) or bool(re.search(r"\b0\d[\d ]{7,}\b", l))


def words(text):
    # Unicode letters, so accented words ("café") are checked whole instead of in fragments.
    return re.findall(r"[^\W\d_](?:[^\W\d_]|[&/+'\-])*[^\W\d_]|[^\W\d_]", text)


def numbers(text):
    return {m.replace(",", "") for m in re.findall(r"\d[\d,]*(?:\.\d+)?", text)}


def is_year(n):
    return re.fullmatch(r"(19|20)\d\d", n) is not None


def supported(word, master_low, master_words):
    wl = word.lower()
    if wl in master_words:
        return True
    pre = wl[:6] if len(wl) >= 6 else wl          # morphology: negotiation~negotiate
    if pre and pre in master_low:
        return True
    if len(wl) >= 4 and wl in master_low:
        return True
    return False


def main():
    if any(a in ("-h", "--help") for a in sys.argv[1:]):
        print(__doc__)
        return
    args = [a for a in sys.argv[1:] if a != "--no-color"]
    if len(args) < 2:
        print("Usage: python check_grounding.py cv.docx master_profile.(md|txt|docx)", file=sys.stderr)
        sys.exit(2)
    cv_path, master_path = args[0], args[1]

    master = read_text(master_path)
    master_low = master.lower()
    master_words = {w.lower() for w in words(master)}
    master_nums = numbers(master)

    raw = [l for l in get_lines(cv_path) if not is_contact(l)]

    # Ignore identity + structural lines: the name (first line), a short target
    # headline (second line, no sentence punctuation), and any section heading.
    lines = []
    for idx, line in enumerate(raw):
        s = line.strip()
        low = s.lower().rstrip(":")
        if low in HEADINGS:
            continue
        if idx == 0:                                   # name
            continue
        if idx == 1 and len(s.split()) <= 6 and "." not in s and "," not in s:
            continue                                   # target headline (e.g. "Head of Operations")
        lines.append(line)

    bad_nums, flagged_lines, distinct = [], [], {}
    for line in lines:
        terms = []
        for w in words(line):
            wl = w.lower()
            if len(wl) < 4 or wl in STOP:
                continue
            if not supported(w, master_low, master_words):
                terms.append(w)
                distinct[wl] = distinct.get(wl, 0) + 1
        for n in numbers(line):
            digits = n.replace(".", "")
            if is_year(n) or ("." not in n and len(digits) < 2):
                continue
            if n not in master_nums:
                bad_nums.append((n, line))
        if terms:
            flagged_lines.append((line, sorted(set(t.lower() for t in terms))))

    print(f"Grounding check\n  CV:     {cv_path}\n  master: {master_path}\n")

    print("== UNSUPPORTED NUMBERS (possible fabricated metrics) ==")
    if bad_nums:
        seen = set()
        for n, line in bad_nums:
            if (n, line) in seen:
                continue
            seen.add((n, line))
            print(f"{RED}FAIL{RESET}  '{n}' is not in the master profile  ->  {line.strip()[:104]}")
    else:
        print(f"{GREEN}PASS{RESET}  Every metric in the CV also appears in the master profile.")

    print("\n== LINES WITH WORDS NOT FOUND IN THE MASTER PROFILE ==")
    if flagged_lines:
        print(f"{YELLOW}REVIEW{RESET}  {len(flagged_lines)} line(s) use words absent from the master profile. "
              "For each, confirm it is a fair restatement of something you genuinely did, and CUT "
              "anything that isn't - this is exactly how invented claims slip onto a CV.")
        for line, terms in flagged_lines:
            print(f"  - {line.strip()[:118]}")
            print(f"      not in master: {', '.join(terms)}")
    else:
        print(f"{GREEN}PASS{RESET}  Every content word in the CV is traceable to the master profile.")

    print(f"\nSummary: {len(bad_nums)} unsupported number(s), {len(flagged_lines)} line(s) to review, "
          f"{len(distinct)} distinct unsupported term(s).")
    print("Note: over-flagging is intentional - legitimate JD-mirrored synonyms appear here too. "
          "Verify each flagged line; remove unsupported claims.")


if __name__ == "__main__":
    main()
